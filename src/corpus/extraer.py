"""Extracción de texto limpio desde HTML, PDF, DOCX y TXT.

`extraer(ruta)` acepta un archivo o una carpeta; en una carpeta concatena los
archivos soportados en orden alfabético (así se unen las páginas `_pr001…` que
la Secretaría del Senado publica por separado).
"""
from __future__ import annotations

import html
import re
import shutil
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

EXTENSIONES = {".html", ".htm", ".pdf", ".docx", ".txt"}

_BLOQUES = ("p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6",
            "table", "section", "article", "blockquote", "pre", "dd", "dt")
_RUIDO = ("script", "style", "noscript", "nav", "header", "footer", "form", "iframe",
          "button", "select", "svg", "head")


@dataclass
class Extraccion:
    texto: str
    archivos: list[Path] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- utilidades

def decodificar(datos: bytes) -> str:
    """UTF-8 si es válido; si no, windows-1252 (codificación habitual del Senado y SUIN)."""
    if datos.startswith(b"\xef\xbb\xbf"):
        datos = datos[3:]
    try:
        return datos.decode("utf-8")
    except UnicodeDecodeError:
        return datos.decode("cp1252", errors="replace")


def normalizar(texto: str) -> str:
    """NFC, espacios uniformes, sin espacios al borde de línea, máximo una línea en blanco."""
    texto = unicodedata.normalize("NFC", texto)
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    texto = re.sub(r"[   \t\f\v]", " ", texto)
    texto = re.sub(r"[​‌‍﻿­]", "", texto)
    lineas = [re.sub(r" {2,}", " ", l).strip() for l in texto.split("\n")]
    texto = "\n".join(lineas)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip() + "\n" if texto.strip() else ""


def _unir_guiones(texto: str) -> str:
    """Une palabras partidas por guion al final de línea (típico de PDF)."""
    return re.sub(r"(\w)-\n(\w)", r"\1\2", texto)


# --------------------------------------------------------------------------- formatos

def extraer_html(datos: bytes) -> str:
    from bs4 import BeautifulSoup

    sopa = BeautifulSoup(decodificar(datos), "html.parser")
    for tag in sopa.find_all(_RUIDO):
        tag.decompose()
    for br in sopa.find_all("br"):
        br.replace_with("\n")
    for tag in sopa.find_all(_BLOQUES):
        tag.insert_before("\n")
        tag.insert_after("\n")
    for celda in sopa.find_all(("td", "th")):
        celda.insert_after(" ")
    cuerpo = sopa.body or sopa
    return normalizar(cuerpo.get_text(""))


def _ocr_disponible() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def extraer_pdf(ruta: Path, avisos: list[str]) -> str:
    import pymupdf

    paginas: list[list[str]] = []
    sin_texto = 0
    with pymupdf.open(ruta) as doc:
        for pagina in doc:
            texto = pagina.get_text("text")
            if not texto.strip():
                sin_texto += 1
                if _ocr_disponible():
                    import pytesseract
                    from PIL import Image

                    pix = pagina.get_pixmap(dpi=300)
                    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    texto = pytesseract.image_to_string(img, lang="spa")
            paginas.append(texto.split("\n"))
    if sin_texto:
        modo = "con OCR" if _ocr_disponible() else "SIN OCR (instale tesseract + pytesseract)"
        avisos.append(f"{ruta.name}: {sin_texto} página(s) sin capa de texto, procesadas {modo}")

    # Encabezados/pies repetidos: líneas cortas presentes en ≥ 60 % de las páginas.
    if len(paginas) >= 3:
        conteo = Counter(l.strip() for p in paginas for l in set(p) if l.strip())
        repetidas = {l for l, n in conteo.items() if n >= 0.6 * len(paginas) and len(l) <= 80}
    else:
        repetidas = set()
    lineas = []
    for p in paginas:
        for l in p:
            s = l.strip()
            if s in repetidas or re.fullmatch(r"(p[áa]g(ina)?\.?\s*)?\d{1,4}(\s*de\s*\d{1,4})?", s,
                                             re.IGNORECASE):
                continue
            lineas.append(l)
        lineas.append("")
    return normalizar(_unir_guiones("\n".join(lineas)))


def _notas_docx(ruta: Path) -> list[str]:
    """Texto de las notas al pie y al final de un .docx (python-docx no las lee).

    En las sentencias estas notas suelen citar normas y precedentes.
    """
    import zipfile

    notas: list[str] = []
    with zipfile.ZipFile(ruta) as z:
        for parte in ("word/footnotes.xml", "word/endnotes.xml"):
            if parte not in z.namelist():
                continue
            xml = z.read(parte).decode("utf-8", errors="ignore")
            for nota in re.findall(r"<w:(?:footnote|endnote)\b(.*?)</w:(?:footnote|endnote)>", xml,
                                   flags=re.DOTALL):
                if re.search(r'w:type="(?:separator|continuationSeparator|continuationNotice)"', nota):
                    continue
                texto = " ".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", nota)).strip()
                if texto:
                    notas.append(html.unescape(texto))
    return notas


def extraer_docx(ruta: Path) -> str:
    import docx

    documento = docx.Document(str(ruta))
    partes = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        for fila in tabla.rows:
            partes.append(" ".join(c.text.strip() for c in fila.cells))
    notas = _notas_docx(ruta)
    if notas:
        partes.append("\nNOTAS")
        partes += [f"[{n}] {t}" for n, t in enumerate(notas, start=1)]
    return normalizar("\n".join(partes))


def extraer_archivo(ruta: Path, avisos: list[str]) -> str:
    ext = ruta.suffix.lower()
    if ext in (".html", ".htm"):
        return extraer_html(ruta.read_bytes())
    if ext == ".pdf":
        return extraer_pdf(ruta, avisos)
    if ext == ".docx":
        return extraer_docx(ruta)
    if ext == ".txt":
        return normalizar(decodificar(ruta.read_bytes()))
    raise ValueError(f"Formato no soportado: {ruta.name}")


def extraer(ruta: Path) -> Extraccion:
    """Texto limpio de un archivo o de todos los archivos soportados de una carpeta."""
    if ruta.is_dir():
        archivos = sorted((p for p in ruta.rglob("*")
                           if p.is_file() and p.suffix.lower() in EXTENSIONES),
                          key=lambda p: p.relative_to(ruta).as_posix().lower())
    else:
        archivos = [ruta]
    res = Extraccion(texto="", archivos=archivos)
    if not archivos:
        res.avisos.append(f"{ruta}: no hay archivos soportados ({sorted(EXTENSIONES)})")
        return res
    partes = []
    for a in archivos:
        try:
            texto = extraer_archivo(a, res.avisos)
        except Exception as exc:  # un archivo dañado no debe tumbar toda la ingesta
            res.avisos.append(f"{a.name}: error al extraer ({exc})")
            continue
        if not texto.strip():
            res.avisos.append(f"{a.name}: no se extrajo texto")
        partes.append(texto)
    res.texto = normalizar("\n\n".join(partes))
    return res
