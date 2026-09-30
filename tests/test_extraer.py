"""Pruebas de los extractores con archivos generados en la prueba (TEXTO DE PRUEBA)."""
from pathlib import Path

from src.corpus.extraer import extraer, normalizar

HTML = """<html><head><meta charset="windows-1252"><title>Titulo</title>
<script>var x = 1;</script><style>p{}</style></head>
<body><nav>Inicio | Menú</nav>
<p>TEXTO DE PRUEBA. ART&Iacute;CULO 1o. Primera&nbsp;disposición de prueba.</p>
<p>ARTÍCULO 2o. Segunda<br>línea de prueba.</p>
<table><tr><td>celda a</td><td>celda b</td></tr></table>
<footer>Pie de página</footer></body></html>"""


def test_html_cp1252_limpia_ruido(tmp_path):
    p = tmp_path / "norma.html"
    p.write_bytes(HTML.encode("cp1252"))
    res = extraer(p)
    t = res.texto
    assert "ARTÍCULO 1o. Primera disposición de prueba." in t
    assert "Segunda\nlínea de prueba." in t
    assert "celda a celda b" in t
    for ruido in ("var x", "Menú", "Pie de página", "Titulo"):
        assert ruido not in t


def test_carpeta_une_paginas_en_orden(tmp_path):
    d = tmp_path / "codigo"
    d.mkdir()
    (d / "codigo_pr002.html").write_text("<p>ARTÍCULO 3o. Tercera.</p>", encoding="utf-8")
    (d / "codigo.html").write_text("<p>ARTÍCULO 1o. Primera.</p>", encoding="utf-8")
    (d / "codigo_pr001.html").write_text("<p>ARTÍCULO 2o. Segunda.</p>", encoding="utf-8")
    (d / "notas.xlsx").write_bytes(b"ignorado")
    t = extraer(d).texto
    assert t.index("Primera") < t.index("Segunda") < t.index("Tercera")


def test_pdf_con_texto_quita_encabezados_y_numeros(tmp_path):
    import pymupdf

    ruta = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    for i in range(3):
        pag = doc.new_page()
        pag.insert_text((72, 40), "DIARIO DE PRUEBA")
        pag.insert_text((72, 100), f"ARTICULO {i + 1}. Texto de prueba numero {i + 1}.")
        pag.insert_text((72, 800), str(i + 1))
    doc.save(ruta)
    doc.close()
    t = extraer(ruta).texto
    assert "ARTICULO 1. Texto de prueba numero 1." in t
    assert "ARTICULO 3." in t
    assert "DIARIO DE PRUEBA" not in t
    assert "\n2\n" not in t


def test_pdf_sin_texto_avisa(tmp_path):
    import pymupdf

    ruta = tmp_path / "escaneado.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(ruta)
    doc.close()
    res = extraer(ruta)
    assert any("sin capa de texto" in a for a in res.avisos)


def test_docx(tmp_path):
    import docx

    ruta = tmp_path / "doc.docx"
    d = docx.Document()
    d.add_paragraph("ARTÍCULO 1. Párrafo de prueba.")
    tabla = d.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "uno"
    tabla.rows[0].cells[1].text = "dos"
    d.save(ruta)
    t = extraer(ruta).texto
    assert "ARTÍCULO 1. Párrafo de prueba." in t and "uno dos" in t


def test_docx_incluye_notas_al_pie(tmp_path):
    import zipfile

    import docx

    base = tmp_path / "base.docx"
    d = docx.Document()
    d.add_paragraph("Texto principal de prueba.")
    d.save(base)
    notas = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
             '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t>---</w:t></w:r></w:p></w:footnote>'
             '<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">Ver Ley 1150 de 2007 </w:t></w:r>'
             '<w:r><w:t>y Sentencia C-355 de 2006.</w:t></w:r></w:p></w:footnote></w:footnotes>')
    ruta = tmp_path / "con_notas.docx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(ruta, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/footnotes.xml", notas)
    t = extraer(ruta).texto
    assert "Texto principal de prueba." in t
    assert "[1] Ver Ley 1150 de 2007 y Sentencia C-355 de 2006." in t
    assert "---" not in t                                  # el separador no es contenido


def test_txt_cp1252(tmp_path):
    ruta = tmp_path / "doc.txt"
    ruta.write_bytes("Artículo 5. Niño y compañía.\r\n\r\n\r\n\r\nFin".encode("cp1252"))
    assert extraer(ruta).texto == "Artículo 5. Niño y compañía.\n\nFin\n"


def test_normalizar():
    assert normalizar("  a  b \n\n\n\n c​ ") == "a b\n\nc\n"
    assert normalizar("   \n  ") == ""


def test_archivo_danado_no_tumba(tmp_path):
    d = tmp_path / "x"
    d.mkdir()
    (d / "a.pdf").write_bytes(b"no es un pdf")
    (d / "b.txt").write_text("contenido", encoding="utf-8")
    res = extraer(d)
    assert "contenido" in res.texto
    assert any("error al extraer" in a for a in res.avisos)
